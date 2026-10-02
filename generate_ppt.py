import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    # Create a presentation object
    prs = Presentation()

    # Define slide layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    
    # -----------------------------------------------------------------
    # Slide 1: Title Slide
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "BitScan Suite"
    subtitle.text = "Advanced Data Sanitization & Forensic Recovery Platform\nBuilt with Python & PyQt6"
    
    # -----------------------------------------------------------------
    # Slide 2: The Problem
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "The Problem: Hidden Digital Evidence"
    tf = body_shape.text_frame
    tf.text = "Deleted files are rarely erased immediately from storage media."
    
    p = tf.add_paragraph()
    p.text = "Operating systems only delete the file pointers, leaving the raw data intact."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Standard recovery tools often rely on corrupted file allocation tables (FAT/NTFS) and fail to find hidden data."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Need for an independent, sector-level tool that scans raw bytes bypassing the OS."
    p.level = 1

    # -----------------------------------------------------------------
    # Slide 3: The BitScan Solution
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "The Solution: BitScan Architecture"
    tf = body_shape.text_frame
    tf.text = "BitScan performs deep forensic carving directly on physical drives."
    
    p = tf.add_paragraph()
    p.text = "Raw Disk I/O Access: Bypasses Windows restrictions to read \\\\.\\PhysicalDrive at the hardware level."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Magic Byte Signatures: Identifies files (JPEG, PDF, ZIP) via hexadecimal file headers rather than OS file tables."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Strict Sector Alignment: Reads data in perfect 512-byte blocks to prevent OS caching errors and hardware exceptions."
    p.level = 1

    # -----------------------------------------------------------------
    # Slide 4: Key Features
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Key Features & Capabilities"
    tf = body_shape.text_frame
    
    tf.text = "Google Chrome Material UI: A world-class, snappy user interface."
    
    p = tf.add_paragraph()
    p.text = "Selective Extraction: Users can review carved artifacts and selectively recover only relevant evidence."
    
    p = tf.add_paragraph()
    p.text = "Tamper-Evident Reporting: Generates SHA-256 hashes for every recovered file to maintain the chain of custody."

    p = tf.add_paragraph()
    p.text = "Device Discovery: Uses PowerShell CIM integration to automatically detect physical and virtual loop devices safely."

    # -----------------------------------------------------------------
    # Slide 5: Why Python?
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Technology Stack: Is Python Powerful Enough?"
    tf = body_shape.text_frame
    tf.text = "Yes. Python serves as an incredible orchestration layer."
    
    p = tf.add_paragraph()
    p.text = "Native C++ UI: The PyQt6 interface is a Python wrapper around the native Qt C++ framework, ensuring flawless rendering speeds."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "I/O Bottlenecks: Disk read speeds (150MB/s - 500MB/s) are the limiting factor, not the Python interpreter."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Rapid Prototyping: Python allows for extreme agility in adapting magic byte signatures and forensic logic on the fly."
    p.level = 1

    # -----------------------------------------------------------------
    # Slide 6: Future Roadmap
    # -----------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Future Roadmap & Scaling"
    tf = body_shape.text_frame
    
    tf.text = "Multiprocessing Architecture"
    p = tf.add_paragraph()
    p.text = "Implement parallel background workers to scan high-speed NVMe drives across multiple CPU cores simultaneously."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Heuristic Validation"
    
    p = tf.add_paragraph()
    p.text = "Analyze internal file structures (not just headers) to automatically discard corrupted artifacts (false positives)."
    p.level = 1

    # Save presentation
    prs.save("BitScan_Presentation.pptx")
    print("Successfully generated BitScan_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
